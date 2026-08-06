from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r'C:\Users\Admin\Desktop\Draft\DraftDatastore_Phase2_Handoff.docx')
BLUE = '2E74B5'; DARK = '1F4D78'; LIGHT = 'E8EEF5'; GREY = 'F2F4F7'; INK = '0B2545'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_width(cell, dxa):
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn('w:tcW'))
    if tcW is None: tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(dxa)); tcW.set(qn('w:type'), 'dxa')

def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW') or OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '9360'); tblW.set(qn('w:type'), 'dxa')
    if tblW.getparent() is None: tblPr.append(tblW)
    ind = OxmlElement('w:tblInd'); ind.set(qn('w:w'), '120'); ind.set(qn('w:type'), 'dxa'); tblPr.append(ind)
    grid = table._tbl.tblGrid
    for gridCol, width in zip(grid.gridCol_lst, widths): gridCol.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
            if mar is None:
                mar = OxmlElement('w:tcMar'); tcPr.append(mar)
            for side in ('top','start','bottom','end'):
                el = OxmlElement(f'w:{side}'); el.set(qn('w:w'), '80' if side in ('top','bottom') else '120'); el.set(qn('w:type'), 'dxa'); mar.append(el)

def style_cell(cell, bold=False, color=None, size=9.5):
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(1)
        for run in paragraph.runs:
            run.font.name='Calibri'; run._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
            run.font.size=Pt(size); run.bold=bold
            if color: run.font.color.rgb=RGBColor.from_string(color)

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers)); table.style='Table Grid'
    set_table_geometry(table, widths)
    for i,h in enumerate(headers):
        cell=table.rows[0].cells[i]; cell.text=h; set_cell_shading(cell, LIGHT); style_cell(cell, True, INK)
    for ridx,row in enumerate(rows):
        cells=table.add_row().cells
        for i,value in enumerate(row):
            cells[i].text=str(value); style_cell(cells[i])
            if ridx % 2 == 1: set_cell_shading(cells[i], 'FAFBFC')
    doc.add_paragraph().paragraph_format.space_after=Pt(3)
    return table

def add_heading(doc, text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.add_run(text); return p

def add_body(doc, text, bold_prefix=None):
    p=doc.add_paragraph(style='Normal')
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold=True; p.add_run(text[len(bold_prefix):])
    else: p.add_run(text)
    return p

def add_numbered(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Number'); p.add_run(item)

doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=Inches(.492); sec.footer_distance=Inches(.492)

styles=doc.styles
normal=styles['Normal']; normal.font.name='Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
for name,size,color,before,after in [('Heading 1',16,BLUE,18,10),('Heading 2',13,BLUE,14,7),('Heading 3',12,DARK,10,5)]:
    st=styles[name]; st.font.name='Calibri'; st._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True; st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; rr=footer.add_run('Draft Datastore | Phase 2 Handoff | 25 July 2026'); rr.font.name='Calibri'; rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string('666666')

title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.LEFT; title.paragraph_format.space_after=Pt(3)
r=title.add_run('Draft Datastore'); r.font.name='Calibri'; r.font.size=Pt(24); r.bold=True; r.font.color.rgb=RGBColor.from_string(INK)
sub=doc.add_paragraph(); sub.paragraph_format.space_after=Pt(14); rr=sub.add_run('Phase 2 Backend Implementation Handoff'); rr.font.name='Calibri'; rr.font.size=Pt(15); rr.font.color.rgb=RGBColor.from_string(DARK)
add_body(doc, 'Purpose: provide enough accurate implementation, operational, and decision context for another Codex agent to resume backend development without rediscovering the work already completed.')
add_table(doc, ['Field','Current state'], [
    ('Workspace','C:\\Users\\Admin\\Desktop\\Draft'),
    ('Backend solution','API\\DraftDatastore\\DraftDatastore.sln'),
    ('Target runtime','.NET 9 / ASP.NET Core 9'),
    ('Completed modules','1 Solution Structure; 2 Database; 3 Authentication'),
    ('Next module','4 Player module'),
    ('Last successful check','dotnet build DraftDatastore.sln --no-restore'),
], [2700,6660])

add_heading(doc,'1. Resume Instructions')
add_numbered(doc,[
    'Set the working directory to C:\\Users\\Admin\\Desktop\\Draft\\API\\DraftDatastore.',
    'Read this handoff, then inspect the current working tree before changing files. The legacy API\\DraftAOD template and UI\\DraftAOD frontend are user-owned and intentionally untouched.',
    'Run dotnet build DraftDatastore.sln --no-restore. If restore is required, NuGet network access may need escalation because sandbox TLS access previously failed.',
    'Continue with Module 4: Player module. Preserve the existing clean-architecture reference direction and thin-controller rule.',
    'Before adding a migration, set DRAFT_DATASTORE_CONNECTION_STRING only for the command session; never commit a production connection string or secret.',
    'After each module, build the entire solution and document the decision rationale, alternatives, interview questions, and scalability factors.'
])

add_heading(doc,'2. Approved Scope and Constraints')
add_body(doc,'Backend only. Do not generate Angular code or Azure deployment configuration at this stage. The approved stack is ASP.NET Core 9, EF Core, SQL Server, JWT, refresh tokens, AutoMapper, FluentValidation, Serilog, Swagger, DI, repositories, and Clean Architecture.')
add_body(doc,'The chatbot must be database-grounded only. No LLM or external football-knowledge source may be called. The required progression is database, authentication, player catalogue, admin, chatbot, logging, validation, middleware, Swagger, then deployment configuration.')

add_heading(doc,'3. Solution Structure (Module 1)')
add_table(doc,['Project','Purpose','References'],[
    ('DraftDatastore.Domain','Entities, role constants, audit and soft-delete contracts','None'),
    ('DraftDatastore.Application','DTOs, validators, interfaces, use-case orchestration','Domain'),
    ('DraftDatastore.Persistence','EF Core DbContext, model configuration, migrations','Domain'),
    ('DraftDatastore.Infrastructure','JWT, password/token services, EF-backed stores, DI','Application, Persistence'),
    ('DraftDatastore.API','Controllers, HTTP pipeline, Swagger, auth configuration','Application, Infrastructure, Persistence'),
],[2100,4500,2760])
add_body(doc,'The legacy default template remains at API\\DraftAOD and has not been converted. All active implementation work is under API\\DraftDatastore.')

add_heading(doc,'4. Database Module (Module 2)')
add_body(doc,'Primary implementation files: DraftDatastore.Domain\\Entities\\Entities.cs; DraftDatastore.Domain\\Common\\AuditableEntity.cs; DraftDatastore.Persistence\\DraftDatastoreDbContext.cs; DraftDatastore.Persistence\\DraftDatastoreDbContextFactory.cs.')
add_table(doc,['Area','Implemented details'],[
    ('Identity','User, Role, UserRole, RefreshToken; normalized unique email; role seeds User and Admin.'),
    ('Catalogue','Player, PlayerAlias, PlayerPosition, PlayerImage, Position, Nationality, PlayingEra, Favorite.'),
    ('Operations','AuditLog, ChatHistory, LoginHistory.'),
    ('Integrity','Positive rank; non-negative points; era start year <= end year; unique aliases, emails, role/position joins, blob paths, and favourites.'),
    ('Soft delete','User and Player implement ISoftDeletable; SaveChanges converts deletes; global query filters hide deleted rows.'),
    ('Indexes','Rank, nationality, points, deleted state, token hash, login history, chat history, and audit lookup indexes.'),
    ('Migrations','InitialCreate and AddLoginHistory are committed under DraftDatastore.Persistence\\Migrations.'),
],[2100,7260])
add_body(doc,'Design-time EF factory requires DRAFT_DATASTORE_CONNECTION_STRING. It deliberately throws when the variable is absent, preventing accidental hard-coded secrets.')

add_heading(doc,'5. Authentication Module (Module 3)')
add_table(doc,['Component','Path / behavior'],[
    ('Contracts and validation','Application\\Authentication\\AuthenticationContracts.cs: RegisterRequest, LoginRequest, RefreshRequest, AuthResponse; FluentValidation rules including 12-character password policy.'),
    ('Service orchestration','Application\\Authentication\\AuthService.cs: registration, login, refresh rotation, logout, login-history writes.'),
    ('Abstractions','Application\\Authentication\\AuthenticationAbstractions.cs: IIdentityStore, IPasswordService, IJwtTokenService.'),
    ('Persistence implementation','Infrastructure\\Authentication\\EfIdentityStore.cs uses DbContext and includes roles when authenticating.'),
    ('Passwords','Infrastructure\\Authentication\\AspNetPasswordService.cs uses ASP.NET PasswordHasher<User>.'),
    ('Tokens','Infrastructure\\Authentication\\JwtTokenService.cs creates 15-minute access tokens, 64-byte random refresh tokens, and SHA-256 token hashes.'),
    ('Endpoints','API\\Controllers\\AuthController.cs exposes register, login, refresh, logout, and me under /api/v1/auth.'),
    ('Pipeline','API request-validation filter and exception middleware produce 400/401/409/500 Problem Details.'),
],[2100,7260])
add_body(doc,'JWT claims: subject, email, name identifier, display name, and one role claim per assigned role. Authorization policy name Admin requires the Admin role.')

add_heading(doc,'6. API Contract Implemented So Far')
add_table(doc,['Method','Route','Access','Success'],[
    ('POST','/api/v1/auth/register','Anonymous','201 Created + AuthResponse'),
    ('POST','/api/v1/auth/login','Anonymous','200 OK + AuthResponse'),
    ('POST','/api/v1/auth/refresh','Anonymous','200 OK + rotated AuthResponse'),
    ('POST','/api/v1/auth/logout','Authenticated','204 No Content'),
    ('GET','/api/v1/auth/me','Authenticated','200 OK + current claims'),
],[1450,3500,1800,2610])
add_body(doc,'Refresh tokens are passed in the request body at present. Secure HttpOnly cookie support remains an approved optional extension; do not add it without coordinating CORS and CSRF protections with the frontend plan.')

add_heading(doc,'7. Configuration, Secrets, and Tooling')
add_table(doc,['Item','Current state / action for next agent'],[
    ('Connection string','Development LocalDB value is in API\\DraftDatastore.API\\appsettings.json. Production must use environment variables/Key Vault later.'),
    ('JWT signing key','appsettings.json contains a development-only placeholder. Override with a 32+ character secret in production; do not treat the committed value as a production secret.'),
    ('NuGet','Initial sandbox downloads failed TLS authentication. Escalated dotnet restore / dotnet add commands succeeded.'),
    ('EF CLI','dotnet-ef 9.0.7 was installed globally during this work. Use dotnet ef migrations add with Persistence as project/startup project and temporary DRAFT_DATASTORE_CONNECTION_STRING.'),
    ('Build','Successful command: dotnet build DraftDatastore.sln --no-restore.'),
],[2100,7260])

add_heading(doc,'8. Known Decisions and Accepted Risks')
add_table(doc,['Decision','Reason / alternative / impact'],[
    ('Clean Architecture with five projects','Chosen to retain domain independence and testable boundaries. A single API project was rejected because it couples controllers, EF, and business logic. Supports replacing infrastructure and scaling stateless API instances.'),
    ('EF Core DbContext as unit of work','Chosen because DbContext already provides transactions and change tracking. A custom Unit of Work was rejected as redundant wrapping.'),
    ('JWT plus rotating refresh tokens','Chosen for scalable stateless API access while enabling revocation and replay detection. Long-lived JWT-only sessions were rejected.'),
    ('Hashed refresh tokens','Chosen so a database leak does not expose reusable credentials. Plaintext storage was rejected.'),
    ('AutoMapper NU1903 suppression','User explicitly selected retention of AutoMapper despite NuGet flagging GHSA-rvv3-g6hj-g44x. Only NU1903 is suppressed in Application csproj; reassess this dependency before production release.'),
    ('Generated migration CA1861 suppression','CA1861 is suppressed only in Persistence because EF-generated seed-data arrays trigger it. Application code remains analyzer-clean.'),
    ('API logger CA1848 suppression','CA1848 is currently suppressed in API because the first exception middleware uses ILogger extension calls. Replace with LoggerMessage source-generated delegates when refining logging module.'),
],[2600,6760])

add_heading(doc,'9. Important Gaps Before Production')
add_numbered(doc,[
    'No Player, Admin, Chatbot, media upload, audit writer, caching, or Postman collection has yet been implemented.',
    'No automated unit or integration tests currently exist. Add tests as modules are completed, especially token rotation, unauthorized access, soft-delete filtering, and player query sorting.',
    'Migrations have been generated but not applied to a verified SQL Server database in this session.',
    'The application does not yet apply migrations at startup; decide deployment migration ownership later rather than silently applying schema changes in production.',
    'Swagger is enabled in Development but JWT bearer security definitions/examples still need configuration.',
    'The exception middleware is present early because auth needs correct 401 responses; expand it during the dedicated middleware/error-handling module.'
])

add_heading(doc,'10. Recommended Module 4 Plan: Player Catalogue')
add_numbered(doc,[
    'Define Player request/response DTOs, paged response model, query filter model, AutoMapper profile, and FluentValidation rules in Application.',
    'Add repository interfaces in Application and EF implementations in Persistence/Infrastructure. Repositories must contain data access only, not business rules.',
    'Implement PlayerService with async create, read, update, rank update, soft delete, restore, and paged search/filter/sort operations.',
    'Add PlayersController with authenticated read endpoints and Admin-only write endpoints. Keep controllers limited to HTTP concerns.',
    'Use whitelist-based sort fields; never dynamically concatenate SQL or property names from clients.',
    'Add request/response Swagger examples, tests, and a concise decision record; build the complete solution before handoff.'
])

add_heading(doc,'11. Files Most Likely to Change Next')
add_table(doc,['Location','Expected change'],[
    ('DraftDatastore.Application\\Players','New DTOs, service interfaces/implementation, validators, mapping profile, pagination/filter contracts.'),
    ('DraftDatastore.Application\\Common','Reusable paged result/query contracts if desired.'),
    ('DraftDatastore.Infrastructure or Persistence','Player repository implementation and any query-specific projections.'),
    ('DraftDatastore.API\\Controllers\\PlayersController.cs','New controller; user endpoints authenticated, mutations Admin-only.'),
    ('DraftDatastore.API\\Program.cs','DI registrations and Swagger security/example refinements as needed.'),
    ('DraftDatastore.Persistence\\Migrations','Only if schema changes are actually required; do not add migrations for DTO/service-only work.'),
],[3500,5860])

add_heading(doc,'12. Handoff Verification Checklist')
add_numbered(doc,[
    'Confirm the current solution root and run the successful no-restore build command.',
    'Review the two migrations and current model snapshot before changing entities.',
    'Preserve the user-approved AutoMapper advisory exception unless the user explicitly changes that decision.',
    'Do not modify the untouched legacy DraftAOD backend or the Angular UI unless explicitly requested.',
    'Keep all APIs asynchronous, controllers thin, request validation outside controllers, and database access behind the established application abstractions.',
    'After completing a module, provide the five requested implementation-decision dimensions: why, alternatives, rejection rationale, interview questions, and scalability impact.'
])

doc.core_properties.title='Draft Datastore - Phase 2 Backend Handoff'
doc.core_properties.subject='Continuation guide for Codex agent'
doc.core_properties.author='Codex'
doc.save(OUT)
print(OUT)
